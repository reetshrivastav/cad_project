"""
Generate the complete single-column conference paper in Markdown, DOCX, and PDF formats.
Authors: Reet Shrivastav, Pranshi Gupta, Anushka Thakur
"""
import os
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PAPER_DIR = os.path.abspath("paper")
FIG_DIR = os.path.join(PAPER_DIR, "figures")
MD_PATH = os.path.join(PAPER_DIR, "Adaptive_Ingress_Admission_Control_Paper.md")
DOCX_PATH = os.path.join(PAPER_DIR, "Adaptive_Ingress_Admission_Control_Paper.docx")
PDF_PATH = os.path.join(PAPER_DIR, "Adaptive_Ingress_Admission_Control_Paper.pdf")

# Color palette: Academic Navy / Slate
COLOR_PRIMARY = RGBColor(26, 54, 93)     # Deep Navy
COLOR_SECONDARY = RGBColor(43, 108, 176) # Slate Blue
COLOR_TEXT = RGBColor(45, 55, 72)        # Charcoal Body Text
COLOR_MUTED = RGBColor(113, 128, 150)    # Cool Grey

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text, title="KEY FINDING"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "EDF2F7")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border styling
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="2B6CB0"/>'
        f'<w:top w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = COLOR_SECONDARY
    
    r_text = p.add_run(text)
    r_text.font.name = "Calibri"
    r_text.font.size = Pt(9.5)
    r_text.font.italic = True
    r_text.font.color.rgb = COLOR_TEXT
    doc.add_paragraph() # Spacing

print("Paper generation framework loaded.")
