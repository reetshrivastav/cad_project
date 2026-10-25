"""
Builds a beautifully styled Word DOCX and PDF copy of the Project Review Guide.
"""
import os
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

GUIDE_DIR = os.path.abspath("review_guide")
MD_PATH = os.path.join(GUIDE_DIR, "PROJECT_REVIEW_GUIDE.md")
DOCX_PATH = os.path.join(GUIDE_DIR, "PROJECT_REVIEW_GUIDE.docx")
PDF_PATH = os.path.join(GUIDE_DIR, "PROJECT_REVIEW_GUIDE.pdf")

COLOR_NAVY = RGBColor(26, 54, 93)     # #1A365D
COLOR_BLUE = RGBColor(43, 108, 176)   # #2B6CB0
COLOR_TEXT = RGBColor(45, 55, 72)     # #2D3748
COLOR_MUTED = RGBColor(113, 128, 150) # #718096

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_table(doc, col_widths, headers, data):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], "1A365D")
        set_cell_margins(hdr_cells[i], top=90, bottom=90, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.runs[0]
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    for row_idx, row_data in enumerate(data):
        row_cells = tbl.rows[row_idx + 1].cells
        bg_color = "FFFFFF" if row_idx % 2 == 0 else "F7FAFC"
        for col_idx, text in enumerate(row_data):
            row_cells[col_idx].text = str(text)
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=70, bottom=70, left=100, right=100)
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.runs[0]
            r.font.name = "Calibri"
            r.font.size = Pt(9)
            r.font.color.rgb = COLOR_TEXT
            
    for row in tbl.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width
    doc.add_paragraph()

def build_docx():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        
        p_hdr = s.header.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_hdr = p_hdr.add_run("Project Review & Viva Preparation Guide")
        r_hdr.font.name = "Calibri"
        r_hdr.font.size = Pt(8.5)
        r_hdr.font.color.rgb = COLOR_MUTED
        
        p_ftr = s.footer.paragraphs[0]
        p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_ftr = p_ftr.add_run("Page ")
        r_ftr.font.name = "Calibri"
        r_ftr.font.size = Pt(9)
        r_ftr.font.color.rgb = COLOR_MUTED

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("🎓 Project Review & Viva Preparation Guide")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_NAVY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(4)
    r_sub = p_sub.add_run("Adaptive Ingress Admission Control for Multi-Tenant LLM Serving")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(12)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_BLUE

    p_auth = doc.add_paragraph()
    p_auth.paragraph_format.space_before = Pt(0)
    p_auth.paragraph_format.space_after = Pt(14)
    r_auth = p_auth.add_run("Authors: Reet Shrivastav, Pranshi Gupta, Anushka Thakur\nDepartment of Computer Science & Engineering")
    r_auth.font.name = "Calibri"
    r_auth.font.size = Pt(10)
    r_auth.font.color.rgb = COLOR_MUTED

    # Read markdown and parse into Word
    with open(MD_PATH, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_table = False
    table_lines = []

    for line in lines:
        sline = line.strip()
        if not sline:
            continue
        if sline.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(sline[2:])
            r.font.name = "Calibri"
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = COLOR_NAVY
        elif sline.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(11)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(sline[3:])
            r.font.name = "Calibri"
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = COLOR_BLUE
        elif sline.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(sline[4:])
            r.font.name = "Calibri"
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = COLOR_TEXT
        elif sline.startswith("|") and sline.endswith("|"):
            # Table row
            continue
        elif sline.startswith("> "):
            # Callout
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            c = tbl.cell(0, 0)
            set_cell_background(c, "F7FAFC")
            set_cell_margins(c, top=100, bottom=100, left=140, right=140)
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(sline[2:].replace("*", ""))
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.italic = True
            r.font.color.rgb = COLOR_TEXT
            doc.add_paragraph()
        elif sline.startswith("```"):
            continue
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            
            # Simple run parsing
            r = p.add_run(sline.replace("**", ""))
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.font.color.rgb = COLOR_TEXT

    doc.save(DOCX_PATH)
    print(f"Saved DOCX to {DOCX_PATH}")

def convert_to_pdf():
    ps1_code = f"""
param([string]$DocxPath, [string]$PdfPath)
$word = New-Object -ComObject Word.Application
$word.Visible = $false
try {{
    $doc = $word.Documents.Open($DocxPath)
    $doc.SaveAs([ref]$PdfPath, [ref]17)
    $doc.Close()
}} finally {{
    $word.Quit()
}}
"""
    ps1_path = os.path.join(GUIDE_DIR, "convert_guide.ps1")
    with open(ps1_path, "w", encoding="utf-8") as f:
        f.write(ps1_code)
        
    res = subprocess.run([
        "powershell", "-NoProfile", "-ExecutionPolicy", "Bypass",
        "-File", ps1_path,
        "-DocxPath", DOCX_PATH,
        "-PdfPath", PDF_PATH
    ], capture_output=True, text=True)
    
    if os.path.exists(ps1_path):
        os.remove(ps1_path)
        
    if os.path.exists(PDF_PATH):
        print(f"Successfully generated PDF: {PDF_PATH} ({os.path.getsize(PDF_PATH)} bytes)")
    else:
        print(f"Error compiling PDF: {res.stderr}")

if __name__ == "__main__":
    build_docx()
    convert_to_pdf()
